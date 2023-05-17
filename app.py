from flask import Flask, render_template, url_for, jsonify, send_file, request
import docx
from docx2pdf import convert
import os
from werkzeug.utils import secure_filename
import pythoncom




ALLOWED_EXTENSIONS = {'pdf'}

path = ''

app = Flask(__name__)

app.config['UPLOAD_FOLDER'] = 'uploads'

def replace_placeholder(doc, placeholder, replacement):
    for p in doc.paragraphs:
        if placeholder in p.text:
            inline = p.runs
            for item in range(len(inline)):
                if placeholder in inline[item].text:
                    inline[item].text = inline[item].text.replace(placeholder,
                                                                  replacement)

@app.route('/')
def hello():
    return render_template('index.html')

@app.route('/generate', methods = [ 'GET', 'POST' ])
def generate():
    pythoncom.CoInitialize()
    global path
    if request.method == 'POST':
        file = request.files['coverletter.docx']
        filename = secure_filename(file.filename)
        file.save(os.path.join(app.root_path, app.config['UPLOAD_FOLDER'], filename))
        doc = docx.Document(os.path.join(app.root_path, app.config['UPLOAD_FOLDER'], filename))
        placeholders = {
        "RecipientName": request.form['recipientName'],
        "RecipientTitle": request.form['recipientTitle'],
        "Position": request.form['position'],
        "CompanyName": request.form['companyName'],
        "CompanyAddress": request.form['companyAddress'],
        "CityProvincePostal": request.form['cityProvincePostal']
        }
        for placeholder in placeholders:
            replace_placeholder(doc, placeholder, placeholders[placeholder])
        path = app.root_path + '/uploads/' + placeholders["CompanyName"] + "_" + placeholders["Position"]
        doc.save(path + ".docx")
        convert(path + ".docx" )
    return render_template("downloadFile.html")
            

@app.route('/download')
def download():
    global path
    pathdoc = path + ".docx"
    os.remove(pathdoc)
    pathpdf = path + ".pdf"
    return send_file(pathpdf, as_attachment=True)


if __name__ == '__main__':
    app.run(debug=True)
