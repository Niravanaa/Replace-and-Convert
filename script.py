import tkinter as tk
from tkinter import filedialog as fd
from tkinter import StringVar
from docx import Document
from docx2pdf import convert
import os
import sys


sys.stderr = open("consoleoutput.log", "w")

placeholders = {
        "RecipientName": "",
        "RecipientTitle": "",
        "Position": "",
        "CompanyName": "",
        "CompanyAddress": "",
        "CityProvincePostal": ""
}

filename = ''

def select_file(label):
    global filename

    filetypes = (
        ('Word Document', '*.docx'),
        ('All files', '*.*')
    )

    filename = fd.askopenfilename(
        title = 'Open your template Word document',
        initialdir = '/',
        filetypes = filetypes
    )

    label.configure(text=filename)

def replace_placeholder(doc, placeholder, replacement):
    for p in doc.paragraphs:
        if placeholder in p.text:
            inline = p.runs
            for item in range(len(inline)):
                if placeholder in inline[item].text:
                    inline[item].text = inline[item].text.replace(placeholder, replacement)


def generate_cover_letter():
    global filename
    print(filename)
    doc = Document(filename)
    global placeholders
    for placeholder in placeholders:
        replace_placeholder(doc, placeholder, str(placeholders[placeholder].get()))
    output_filename = f"{str(placeholders['CompanyName'].get())}_{str(placeholders['Position'].get())}.docx"
    doc.save(output_filename)
    convert(output_filename)
    os.remove(output_filename)


def quit_application():
    root.destroy()


# Create the GUI
root = tk.Tk()
root.title("Cover Letter Generator")
root.geometry("400x450")
root.maxsize(400, 450)

# Add choose file button
fileLabel = tk.Label(root, text='', wraplength=350)
choose_file = tk.Button(root, text="Choose a File", command= lambda: select_file(fileLabel))
choose_file.pack(pady=5)
fileLabel.pack(pady=5)

# Add labels and entry fields for each placeholder
for placeholder in placeholders:
    label = tk.Label(root, text=placeholder)
    label.pack(pady=5)
    entry = tk.Entry(root)
    entry.pack()
    placeholders[placeholder] = entry

# Add generate and quit buttons
generate_button = tk.Button(root, text="Generate Cover Letter", command=generate_cover_letter)
generate_button.pack(pady=5)

quit_button = tk.Button(root, text="Quit", command=quit_application)
quit_button.pack(pady=5)

# Start the main loop
root.mainloop()
